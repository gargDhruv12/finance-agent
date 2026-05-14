from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "Finance_Credit_Follow_Up_Email_Agent_Project_Report.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(table) -> None:
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side in ("top", "bottom", "start", "end"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), "120" if side in {"start", "end"} else "80")
        node.set(qn("w:type"), "dxa")


def set_table_width(table) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table)
    set_cell_margins(table)

    header_cells = table.rows[0].cells
    for idx, text in enumerate(headers):
        cell = header_cells[idx]
        cell.width = Inches(widths[idx])
        set_cell_shading(cell, LIGHT_GRAY)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = INK
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].width = Inches(widths[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = cells[idx].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            run = para.add_run(text)
            run.font.size = Pt(9.2)
    document.add_paragraph()
    return table


def add_callout(document: Document, title: str, body: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table)
    set_cell_margins(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(3)
    run = para.add_run(title)
    run.bold = True
    run.font.color.rgb = DARK_BLUE
    run.font.size = Pt(10.5)
    body_para = cell.add_paragraph()
    body_para.paragraph_format.space_after = Pt(0)
    body_run = body_para.add_run(body)
    body_run.font.size = Pt(10)
    document.add_paragraph()


def add_bullet(document: Document, text: str) -> None:
    para = document.add_paragraph(style="List Bullet")
    para.paragraph_format.left_indent = Inches(0.5)
    para.paragraph_format.first_line_indent = Inches(-0.25)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.167
    para.add_run(text)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = INK
    title.paragraph_format.space_after = Pt(6)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = DARK_BLUE

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.color.rgb = BLUE
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.color.rgb = BLUE
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.color.rgb = DARK_BLUE
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "Finance Credit Follow-Up Email Agent"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = DARK_BLUE

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Project Report | AI Enablement Internship")
    footer.runs[0].font.size = Pt(9)


def build() -> None:
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Finance Credit Follow-Up Email Agent")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("AI Enablement Internship Project Report")

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Dhruv | IT Branch, NIT Kurukshetra | Roll No. 123103032")

    add_callout(
        document,
        "Executive Summary",
        "I built this prototype to reduce the manual effort involved in following up on overdue invoices. "
        "The agent reads invoice data, identifies which accounts need attention, drafts context-aware follow-up "
        "emails, and keeps a clear audit trail of every decision. I kept the system safe by making dry-run mode "
        "the default, adding an approval layer before real sends, and separating deterministic finance logic from "
        "LLM-based wording.",
    )

    document.add_heading("1. Project Context", level=1)
    document.add_paragraph(
        "Payment follow-up is a small workflow on paper, but in practice it creates repeated work for finance "
        "teams. Every overdue invoice needs the right tone: too soft, and payment may be delayed; too aggressive, "
        "and the client relationship can suffer. I approached this as an automation problem where the system should "
        "handle routine drafting and logging, while humans remain in control of sensitive decisions."
    )
    document.add_paragraph(
        "The final prototype is a finance follow-up agent that works on structured invoice records. It can run from "
        "the command line or through a Streamlit dashboard, generate email drafts for different overdue stages, and "
        "flag accounts that should be reviewed manually instead of being emailed again."
    )

    document.add_heading("2. What I Set Out To Build", level=1)
    for item in [
        "A working end-to-end flow that starts from invoice data and ends with reviewable email outputs.",
        "A clear escalation policy so email tone changes based on how late the payment is.",
        "Personalized drafts that always include invoice number, client name, amount, due date, days overdue, and payment link.",
        "A safe testing setup where emails are logged instead of sent by default.",
        "An approval path for real sending, so automation does not bypass finance review.",
        "A clean repository with documentation, sample data, sample outputs, and a recruiter-friendly demo path.",
    ]:
        add_bullet(document, item)

    document.add_heading("3. Design Approach", level=1)
    document.add_paragraph(
        "I deliberately avoided making the LLM responsible for financial decisions. The system calculates days overdue "
        "and chooses the escalation stage through deterministic code. The LLM is used only where it is useful: turning "
        "validated invoice facts into a polished email. This makes the prototype easier to audit and safer to extend."
    )
    add_table(
        document,
        ["Step", "Component", "Purpose"],
        [
            ["1", "Data ingestion", "Read invoice rows from CSV or Excel and validate the schema."],
            ["2", "Overdue calculation", "Compute days overdue from the due date using a reproducible demo date when needed."],
            ["3", "Escalation logic", "Select the communication stage using deterministic finance rules."],
            ["4", "Draft generation", "Use Gemini when quota is available, with a reliable template fallback."],
            ["5", "Validation", "Check that required invoice facts are present in the generated draft."],
            ["6", "Approval / dry-run", "Queue real sends for approval and keep demo runs in dry-run mode."],
            ["7", "Auditability", "Write JSON, CSV, SQLite, and local trace records for review."],
        ],
        [0.55, 1.7, 4.1],
    )

    document.add_heading("4. Technical Stack", level=1)
    add_table(
        document,
        ["Layer", "Technology", "Reason"],
        [
            ["Language", "Python", "Simple, readable, and suitable for data processing workflows."],
            ["LLM", "Gemini API", "Chosen as a free-tier-friendly option; model is configurable from .env."],
            ["Framework", "LangGraph-capable workflow", "Useful for representing the agent as clear workflow nodes."],
            ["Validation", "Pydantic", "Structured models for invoice records, decisions, drafts, and audits."],
            ["Data", "pandas CSV/Excel", "Common finance-friendly input format."],
            ["Sending", "Dry-run + SMTP option", "Safe by default, but extensible for approved real sending."],
            ["Observability", "Local tracing + LangSmith", "Local trace file by default; LangSmith can be enabled with credentials."],
            ["UI", "Streamlit", "Dashboard for running, filtering, reviewing, and approving drafts."],
        ],
        [1.15, 1.75, 3.45],
    )

    document.add_heading("5. Follow-Up Policy", level=1)
    add_table(
        document,
        ["Stage", "Trigger", "Tone", "Action"],
        [
            ["1", "1-7 days late", "Warm & Friendly", "A light reminder that assumes the delay may be accidental."],
            ["2", "8-14 days late", "Polite but Firm", "A clearer request for payment status or payment date."],
            ["3", "15-21 days late", "Formal & Serious", "A more direct message asking for response within 48 hours."],
            ["4", "22-30 days late", "Stern & Urgent", "A final notice before the matter is escalated internally."],
            ["5", "More than 30 days late", "Manual Review", "No automated email; finance/legal review is required."],
        ],
        [0.7, 1.5, 1.45, 2.9],
    )

    document.add_heading("6. Implementation Details", level=1)
    for item in [
        "Repository files are placed at the root so GitHub opens directly to README, source code, docs, data, and tests.",
        "The CLI entry point is python -m src.main --today 2026-05-14, which gives a repeatable demo run.",
        "The sample dataset covers all important paths: four follow-up stages, one manual escalation, and one not-overdue invoice.",
        "DRY_RUN=true is the default, so generated messages are visible in logs without sending real emails.",
        "Prompt files are kept out of Git, while the prompt design approach is documented safely.",
        "Gemini calls are cached in SQLite and fall back to deterministic templates if API quota is unavailable.",
        "The Streamlit dashboard supports filters, draft review, approval/rejection, and local trace inspection.",
    ]:
        add_bullet(document, item)

    document.add_heading("7. Demo Results", level=1)
    document.add_paragraph(
        "The sample run processed six invoice records and demonstrated the main behavior of the agent. "
        "Four invoices produced dry-run email drafts, one invoice crossed the manual review threshold, and one "
        "invoice was skipped because it was not overdue."
    )
    add_table(
        document,
        ["Invoice", "Days Overdue", "Stage", "Status", "Generation"],
        [
            ["INV-2026-001", "4", "1", "DRY_RUN", "Draft generated and logged"],
            ["INV-2026-002", "11", "2", "DRY_RUN", "Draft generated and logged"],
            ["INV-2026-003", "18", "3", "DRY_RUN", "Draft generated and logged"],
            ["INV-2026-004", "25", "4", "DRY_RUN", "Draft generated and logged"],
            ["INV-2026-005", "34", "5", "ESCALATED", "No email; manual review required"],
            ["INV-2026-006", "0", "0", "SKIPPED", "No email; invoice is not overdue"],
        ],
        [1.25, 1.0, 0.65, 1.0, 2.6],
    )

    document.add_heading("8. Security And Risk Mitigation", level=1)
    add_table(
        document,
        ["Risk", "Mitigation"],
        [
            ["Prompt injection", "Validated invoice fields, constrained prompts, and post-generation checks."],
            ["Data privacy / PII", "Local processing, fake sample data, minimized fields in prompts, and local logs."],
            ["API key exposure", ".env is ignored by Git; .env.example contains placeholders only."],
            ["Hallucination", "Invoice facts come from validated data; generated email content is checked for mandatory fields."],
            ["Unauthorized access", "Prototype runs locally; production deployment should add authentication and rate limits."],
            ["Email spoofing", "Real sending disabled; production requires verified sender domain, SPF, DKIM, and DMARC."],
            ["Accidental sending", "DRY_RUN=true by default; real sends require SMTP settings and approval."],
        ],
        [1.65, 4.8],
    )

    document.add_heading("9. Implemented Enhancements", level=1)
    for item in [
        "Human approval queue before real email sending.",
        "Optional SMTP sender guarded by dry-run and approval settings.",
        "APScheduler support for recurring invoice scans.",
        "SQLite LLM response caching to reduce repeated Gemini calls.",
        "Local tracing plus optional LangSmith hooks.",
        "Expanded dashboard with filters, approval actions, generated email review, and trace visibility.",
    ]:
        add_bullet(document, item)

    document.add_heading("10. Next Steps", level=1)
    for item in [
        "Connect the agent to a real finance data source.",
        "Add user login if the dashboard is shared with a team.",
        "Use a verified email domain before enabling real client emails.",
        "Mask sensitive client data before sending traces to any external tool.",
    ]:
        add_bullet(document, item)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
